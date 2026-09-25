"""Funções auxiliares: gravação de arquivos, datas e sorteios."""

import json
import math
import random
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def intervalo_datas(inicio: date, fim: date):
    d = inicio
    while d <= fim:
        yield d
        d += timedelta(days=1)


def rng_para(*chaves) -> random.Random:
    """Gerador aleatório determinístico para uma combinação de chaves.

    Garante que o mesmo registro receba sempre o mesmo valor/falha, não importa
    quantas vezes o programa rode (as fontes de referência são republicadas a cada execução).
    """
    return random.Random("|".join(str(c) for c in chaves))


def arredondar_preco(valor: float, passo: float = 0.50) -> float:
    """Arredonda para cima no múltiplo de `passo` (como a cantina faz nos reajustes)."""
    return round(math.ceil(round(valor / passo, 6)) * passo, 2)


def sortear_indices(rng, total: int, taxa: float, minimo: int = 1) -> set[int]:
    """Escolhe quais registros receberão falhas.

    Garante pelo menos `minimo` falhas em fontes pequenas (ex.: cardápio),
    a menos que a taxa seja zero.
    """
    if taxa <= 0 or total == 0:
        return set()
    qtd = min(total, max(minimo, round(total * taxa)))
    return set(rng.sample(range(total), qtd))


def escolha_ponderada(rng, opcoes: dict):
    return rng.choices(list(opcoes), weights=list(opcoes.values()), k=1)[0]


def _formatar(ws, colunas, primeira_linha: int):
    for i, col in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(col) + 4)
        for (cel,) in ws.iter_rows(min_row=primeira_linha, min_col=i, max_col=i):
            if isinstance(cel.value, datetime):
                cel.number_format = "dd/mm/yyyy hh:mm:ss"
            elif isinstance(cel.value, date):
                cel.number_format = "dd/mm/yyyy"


def salvar_excel(caminho: Path, colunas: list[str], linhas: list[list], aba: str = "dados"):
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = aba
    ws.append(colunas)
    for cel in ws[1]:
        cel.font = Font(bold=True)
    for linha in linhas:
        ws.append(linha)
    _formatar(ws, colunas, 2)
    wb.save(caminho)


def ler_excel(caminho: Path) -> tuple[list[str], list[list]]:
    ws = load_workbook(caminho).active
    linhas = [list(l) for l in ws.iter_rows(values_only=True)]
    return linhas[0], linhas[1:]


def anexar_excel(caminho: Path, colunas: list[str], linhas: list[list], aba: str = "dados"):
    """Acrescenta linhas ao final de uma planilha, criando-a se ainda não existir."""
    if not caminho.exists():
        salvar_excel(caminho, colunas, linhas, aba)
        return
    wb = load_workbook(caminho)
    ws = wb.active
    primeira = ws.max_row + 1
    for linha in linhas:
        ws.append(linha)
    _formatar(ws, colunas, primeira)
    wb.save(caminho)


def consolidar_lote(arquivo_lote: Path, historico: Path, aba: str, manter_lote: bool = False) -> int:
    """Acrescenta o conteúdo do arquivo do lote à planilha de histórico e apaga o lote."""
    colunas, linhas = ler_excel(arquivo_lote)
    anexar_excel(historico, colunas, linhas, aba)
    if not manter_lote:
        arquivo_lote.unlink()
    return len(linhas)


def salvar_json(caminho: Path, conteudo):
    caminho.parent.mkdir(parents=True, exist_ok=True)
    with open(caminho, "w", encoding="utf-8") as f:
        json.dump(conteudo, f, ensure_ascii=False, indent=2, default=str)
